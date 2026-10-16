# analytics/tests/test_reports_xlsx.py
"""
XLSX report exports — the parallel export built from the same row
generators the CSV export uses (analytics.reports, via
analytics.xlsx_reports.build_workbook).

Covers: every endpoint returns a loadable workbook, the header row and
formatting match report_specs, totals are real SUM() formulas (blank for
ratio/FCR columns), free-text cells can never be read as a formula by
Excel, permissions/farm-isolation hold exactly as for the CSV siblings,
and an empty period produces a message row instead of an error.
"""
import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from partners.models import FarmPartnerLink
from production.models import (
    Batch,
    DailyRecord,
    FeedDelivery,
    Harvest,
    House,
    InventoryItem,
    InventoryUsageLog,
    SaleEvent,
    TaskCompletion,
    TaskTemplate,
)

pytestmark = pytest.mark.django_db

XLSX_REPORT_NAMES = [
    "report-daily-records-xlsx",
    "report-mortality-summary-xlsx",
    "report-fcr-xlsx",
    "report-feed-margin-xlsx",
    "report-inventory-usage-xlsx",
    "report-routine-completion-xlsx",
    "report-sales-history-xlsx",
]

EXPECTED_HEADER = {
    "report-daily-records-xlsx": [
        "Date", "Batch code", "House", "Mortality – disease",
        "Mortality – heat", "Mortality – culled", "Mortality – unknown",
        "Mortality – total", "Feed used (kg)", "Recorded by", "Corrected?",
    ],
    "report-mortality-summary-xlsx": [
        "Batch code", "House", "Status", "Birds placed", "Birds lost",
        "Mortality (%)", "Lost – disease", "Lost – heat", "Lost – culled",
        "Lost – unknown",
    ],
    "report-fcr-xlsx": [
        "Batch code", "House", "Harvest date", "Cycle (days)",
        "Birds placed", "Birds harvested", "Survival (%)",
        "Total feed (kg)", "Total weight (kg)", "Avg bird weight (kg)", "FCR",
    ],
    "report-feed-margin-xlsx": [
        "Batch code", "House", "Harvest date", "Total weight (kg)",
        "Revenue (PHP)", "Revenue per kg (PHP)", "Allocated feed cost (PHP)",
        "Feed margin (PHP)", "Feed margin per kg (PHP)", "Buyer",
    ],
    "report-inventory-usage-xlsx": [
        "Date", "Item", "Quantity used", "Unit", "Recorded by",
    ],
    "report-routine-completion-xlsx": [
        "Date", "Task", "House", "Completed", "Completed by",
    ],
    "report-sales-history-xlsx": [
        "Date", "Batch code", "House", "Sale type", "Birds sold",
        "Total weight (kg)", "Sale value (PHP)", "Buyer",
    ],
}

HEADER_ROW = 6
FIRST_DATA_ROW = 7


def _url(name, farm):
    return reverse(f"analytics:{name}", kwargs={"farm_pk": farm.pk})


def _workbook(response):
    assert response.status_code == 200
    return load_workbook(io.BytesIO(response.content))


@pytest.fixture
def populated_farm(staffed_farm, owner, worker):
    """Mirrors analytics/tests/test_reports.py's fixture, plus a buyer
    business name containing a formula-injection attempt."""
    farm = staffed_farm
    house1 = House.objects.create(farm=farm, name="House 1", capacity=10000)
    house2 = House.objects.create(farm=farm, name="House 2", capacity=10000)

    harvested = Batch.objects.create(
        house=house1,
        batch_code="B-2026-01",
        initial_bird_count=1000,
        start_date=date(2026, 1, 1),
        created_by=owner,
    )
    for day in range(2, 7):
        DailyRecord.objects.create(
            batch=harvested,
            record_date=date(2026, 1, day),
            mortality_disease=2,
            mortality_heat=1,
            feed_kg=Decimal("120.00"),
            recorded_by=worker,
        )
    buyer_link = FarmPartnerLink.objects.create(
        farm=farm,
        partner=owner,
        link_type=FarmPartnerLink.LinkType.CONSUMER,
        business_name="=1+1",  # formula-injection attempt
        linked_by=owner,
    )
    Harvest.objects.create(
        batch=harvested,
        harvest_date=date(2026, 2, 20),
        birds_harvested=985,
        total_weight_kg=Decimal("1800.00"),
        revenue=Decimal("250000.00"),
        buyer_link=buyer_link,
        recorded_by=owner,
    )
    harvested.status = Batch.Status.HARVESTED
    harvested.save(update_fields=["status"])

    SaleEvent.objects.create(
        batch=harvested,
        sale_date=date(2026, 2, 18),
        sale_type=SaleEvent.SaleType.DRESSED,
        bird_count=985,
        total_weight_kg=Decimal("1800.00"),
        revenue=Decimal("250000.00"),
    )

    active = Batch.objects.create(
        house=house2,
        batch_code="B-2026-02",
        initial_bird_count=800,
        start_date=date(2026, 3, 1),
        created_by=owner,
    )
    DailyRecord.objects.create(
        batch=active,
        record_date=date(2026, 3, 2),
        mortality_unknown=3,
        feed_kg=Decimal("60.00"),
        recorded_by=worker,
    )
    SaleEvent.objects.create(
        batch=active,
        sale_date=date(2026, 3, 5),
        sale_type=SaleEvent.SaleType.LIVE,
        bird_count=50,
        total_weight_kg=Decimal("90.00"),
        revenue=None,
    )

    FeedDelivery.objects.create(
        farm=farm,
        delivery_date=date(2026, 1, 1),
        feed_type=FeedDelivery.FeedType.STARTER,
        quantity_kg=Decimal("2000"),
        unit_cost=Decimal("35.00"),
        recorded_by=owner,
    )

    item = InventoryItem.objects.create(
        farm=farm, name="Disinfectant", unit="litre"
    )
    InventoryUsageLog.objects.create(
        item=item,
        quantity_used=Decimal("5.00"),
        usage_date=date(2026, 1, 10),
        recorded_by=worker,
    )
    InventoryUsageLog.objects.create(
        item=item,
        quantity_used=Decimal("3.00"),
        usage_date=date(2026, 2, 10),
        recorded_by=worker,
    )

    template = TaskTemplate.objects.create(
        farm=farm, name="Morning feed", created_by=owner
    )
    TaskCompletion.objects.create(
        template=template,
        house=house1,
        completion_date=date(2026, 1, 10),
        recorded_by=worker,
    )
    return farm


# ─────────────────────────────────────────────────────────────
# Access — same wall as the CSV siblings
# ─────────────────────────────────────────────────────────────


@pytest.mark.parametrize("name", XLSX_REPORT_NAMES)
class TestXLSXAccess:
    def test_owner_can_download(self, auth, owner, populated_farm, name):
        response = auth(owner).get(_url(name, populated_farm))
        assert response.status_code == 200
        assert response["Content-Type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        disposition = response["Content-Disposition"]
        assert disposition.startswith('attachment; filename="checkngo_')
        assert disposition.endswith('.xlsx"')

    def test_manager_can_download(self, auth, manager, populated_farm, name):
        response = auth(manager).get(_url(name, populated_farm))
        assert response.status_code == 200

    def test_worker_is_forbidden(self, auth, worker, populated_farm, name):
        assert auth(worker).get(_url(name, populated_farm)).status_code == 403

    def test_consumer_is_forbidden(self, auth, consumer, populated_farm, name):
        """CONSUMER is the closest thing this app has to a distinct
        'buyer' role — not a FarmMembership role at all, so it is not
        even a farm member and must be blocked like any outsider."""
        assert auth(consumer).get(_url(name, populated_farm)).status_code == 403

    def test_other_farms_member_is_forbidden(
        self, auth, rival_owner, populated_farm, name
    ):
        assert (
            auth(rival_owner).get(_url(name, populated_farm)).status_code == 403
        )

    def test_anonymous_is_rejected(self, api, populated_farm, name):
        assert api.get(_url(name, populated_farm)).status_code in (401, 403)


# ─────────────────────────────────────────────────────────────
# Shape — every report loads and carries the expected header
# ─────────────────────────────────────────────────────────────


@pytest.mark.parametrize("name", XLSX_REPORT_NAMES)
class TestXLSXShape:
    def test_loads_and_header_matches(self, auth, owner, populated_farm, name):
        wb = _workbook(auth(owner).get(_url(name, populated_farm)))
        ws = wb.active
        header = [
            ws.cell(row=HEADER_ROW, column=c).value
            for c in range(1, len(EXPECTED_HEADER[name]) + 1)
        ]
        assert header == EXPECTED_HEADER[name]

    def test_title_and_farm_name_present(self, auth, owner, populated_farm, name):
        wb = _workbook(auth(owner).get(_url(name, populated_farm)))
        ws = wb.active
        assert "CheckN Go" in ws.cell(row=1, column=1).value
        assert ws.cell(row=2, column=1).value == populated_farm.name

    def test_freeze_panes_and_sheet_name(self, auth, owner, populated_farm, name):
        wb = _workbook(auth(owner).get(_url(name, populated_farm)))
        ws = wb.active
        assert ws.freeze_panes == "A7"
        assert len(ws.title) <= 31


# ─────────────────────────────────────────────────────────────
# Content — formats, real dates, totals formulas, formula-injection guard
# ─────────────────────────────────────────────────────────────


class TestXLSXContent:
    def test_daily_records_date_and_kg_formatting(
        self, auth, owner, populated_farm
    ):
        wb = _workbook(
            auth(owner).get(_url("report-daily-records-xlsx", populated_farm))
        )
        ws = wb.active
        date_cell = ws.cell(row=FIRST_DATA_ROW, column=1)
        assert isinstance(date_cell.value, (date, datetime))
        assert date_cell.number_format == "yyyy-mm-dd"

        feed_col = EXPECTED_HEADER["report-daily-records-xlsx"].index(
            "Feed used (kg)"
        ) + 1
        feed_cell = ws.cell(row=FIRST_DATA_ROW, column=feed_col)
        assert feed_cell.number_format == "#,##0.00"
        assert feed_cell.value == 120.0

    def test_daily_records_totals_row_sums_counts_and_kg(
        self, auth, owner, populated_farm
    ):
        wb = _workbook(
            auth(owner).get(_url("report-daily-records-xlsx", populated_farm))
        )
        ws = wb.active
        n_data = 6  # 5 harvested-batch days + 1 active-batch day
        totals_row = FIRST_DATA_ROW + n_data
        assert ws.cell(row=totals_row, column=1).value == "Total"

        total_col = EXPECTED_HEADER["report-daily-records-xlsx"].index(
            "Mortality – total"
        ) + 1
        formula = ws.cell(row=totals_row, column=total_col).value
        assert formula.startswith("=SUM(")
        assert f"{FIRST_DATA_ROW}" in formula
        assert f"{totals_row - 1}" in formula

    def test_mortality_summary_percent_is_fraction_with_percent_format(
        self, auth, owner, populated_farm
    ):
        wb = _workbook(
            auth(owner).get(_url("report-mortality-summary-xlsx", populated_farm))
        )
        ws = wb.active
        pct_col = EXPECTED_HEADER["report-mortality-summary-xlsx"].index(
            "Mortality (%)"
        ) + 1
        # Row order is batch.start_date — B-2026-01 is first.
        cell = ws.cell(row=FIRST_DATA_ROW, column=pct_col)
        assert cell.number_format == "0.00%"
        # CSV shows "1.50" (meaning 1.50%); the xlsx cell stores the
        # fraction so Excel's % format adds the sign itself.
        assert cell.value == pytest.approx(0.015)

    def test_mortality_summary_totals_blank_for_percent_column(
        self, auth, owner, populated_farm
    ):
        wb = _workbook(
            auth(owner).get(_url("report-mortality-summary-xlsx", populated_farm))
        )
        ws = wb.active
        pct_col = EXPECTED_HEADER["report-mortality-summary-xlsx"].index(
            "Mortality (%)"
        ) + 1
        totals_row = FIRST_DATA_ROW + 2  # 2 batches
        assert ws.cell(row=totals_row, column=1).value == "Total"
        assert ws.cell(row=totals_row, column=pct_col).value is None

        count_col = EXPECTED_HEADER["report-mortality-summary-xlsx"].index(
            "Birds placed"
        ) + 1
        formula = ws.cell(row=totals_row, column=count_col).value
        assert formula.startswith("=SUM(")

    def test_fcr_value_and_totals_blank(self, auth, owner, populated_farm):
        wb = _workbook(
            auth(owner).get(_url("report-fcr-xlsx", populated_farm))
        )
        ws = wb.active
        fcr_col = EXPECTED_HEADER["report-fcr-xlsx"].index("FCR") + 1
        cell = ws.cell(row=FIRST_DATA_ROW, column=fcr_col)
        assert cell.number_format == "0.00"
        assert cell.value == pytest.approx(0.333)

        totals_row = FIRST_DATA_ROW + 1  # 1 harvested batch
        assert ws.cell(row=totals_row, column=fcr_col).value is None

    def test_feed_margin_money_formatting_and_totals(
        self, auth, owner, populated_farm
    ):
        wb = _workbook(
            auth(owner).get(_url("report-feed-margin-xlsx", populated_farm))
        )
        ws = wb.active
        header = EXPECTED_HEADER["report-feed-margin-xlsx"]
        margin_col = header.index("Feed margin (PHP)") + 1
        cell = ws.cell(row=FIRST_DATA_ROW, column=margin_col)
        assert cell.number_format == "#,##0.00"
        assert cell.value == pytest.approx(229000.00)

        totals_row = FIRST_DATA_ROW + 1  # 1 harvested batch
        formula = ws.cell(row=totals_row, column=margin_col).value
        assert formula.startswith("=SUM(")

    def test_feed_margin_buyer_formula_injection_is_defused(
        self, auth, owner, populated_farm
    ):
        """buyer business_name is '=1+1' — it must land in the cell as a
        literal string, never as something Excel would compute."""
        wb = _workbook(
            auth(owner).get(_url("report-feed-margin-xlsx", populated_farm))
        )
        ws = wb.active
        buyer_col = EXPECTED_HEADER["report-feed-margin-xlsx"].index("Buyer") + 1
        cell = ws.cell(row=FIRST_DATA_ROW, column=buyer_col)
        assert cell.value == "=1+1"
        assert cell.data_type == "s"

    def test_sales_history_buyer_formula_injection_is_defused(
        self, auth, owner, populated_farm
    ):
        wb = _workbook(
            auth(owner).get(_url("report-sales-history-xlsx", populated_farm))
        )
        ws = wb.active
        header = EXPECTED_HEADER["report-sales-history-xlsx"]
        buyer_col = header.index("Buyer") + 1
        batch_col = header.index("Batch code") + 1
        rows = [
            r
            for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 2)
            if ws.cell(row=r, column=batch_col).value == "B-2026-01"
        ]
        assert len(rows) == 1
        cell = ws.cell(row=rows[0], column=buyer_col)
        assert cell.value == "=1+1"
        assert cell.data_type == "s"

    def test_sales_history_sale_value_money_format(
        self, auth, owner, populated_farm
    ):
        wb = _workbook(
            auth(owner).get(_url("report-sales-history-xlsx", populated_farm))
        )
        ws = wb.active
        header = EXPECTED_HEADER["report-sales-history-xlsx"]
        value_col = header.index("Sale value (PHP)") + 1
        batch_col = header.index("Batch code") + 1
        for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 2):
            if ws.cell(row=r, column=batch_col).value == "B-2026-01":
                cell = ws.cell(row=r, column=value_col)
                assert cell.number_format == "#,##0.00"
                assert cell.value == pytest.approx(250000.00)

    def test_routine_completion_is_text_grid(self, auth, owner, populated_farm):
        wb = _workbook(
            auth(owner).get(
                _url("report-routine-completion-xlsx", populated_farm),
                {"from": "2026-01-10", "to": "2026-01-10"},
            )
        )
        ws = wb.active
        header = EXPECTED_HEADER["report-routine-completion-xlsx"]
        completed_col = header.index("Completed") + 1
        values = {
            ws.cell(row=r, column=completed_col).value
            for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 2)
        }
        assert values == {"yes", "no"}


# ─────────────────────────────────────────────────────────────
# Empty period — a message row, not an error
# ─────────────────────────────────────────────────────────────


class TestXLSXEmptyFarm:
    @pytest.mark.parametrize("name", XLSX_REPORT_NAMES)
    def test_no_records_message(self, auth, owner, staffed_farm, name):
        response = auth(owner).get(_url(name, staffed_farm))
        assert response.status_code == 200
        wb = _workbook(response)
        ws = wb.active
        assert ws.cell(row=FIRST_DATA_ROW, column=1).value == (
            "No records for this period"
        )
